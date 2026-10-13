from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from build_portable import make_source_zip
from create_demo_data import create_demo_data


PROJECT = Path(__file__).resolve().parents[1]


class PortableSourcePackageTests(unittest.TestCase):
    def test_source_zip_contains_every_runtime_module_and_synthetic_examples(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir)
            examples = output / "ExampleData"
            create_demo_data(examples)
            package = make_source_zip(output, examples)

            with zipfile.ZipFile(package) as archive:
                members = set(archive.namelist())
                required = {
                    f"ErrorBandStudio/{file.relative_to(PROJECT).as_posix()}"
                    for file in (PROJECT / "skin_analysis").glob("*.py")
                }
                required.update({
                    "ErrorBandStudio/main.py",
                    "ErrorBandStudio/run_macos.command",
                    "ErrorBandStudio/run_windows.bat",
                    "ErrorBandStudio/build_portable.py",
                    "ErrorBandStudio/requirements-error-band.txt",
                    "ErrorBandStudio/skin_analysis/csv_safety.py",
                })
                self.assertTrue(required <= members, required - members)
                xlsx = sorted(member for member in members if member.endswith(".xlsx"))
                self.assertEqual(len(xlsx), 20)
                self.assertIsNone(archive.testzip())
                self.assertFalse(any(
                    forbidden in member
                    for member in members
                    for forbidden in (".venv", "__pycache__", ".DS_Store", "site-packages", "/tests/")
                ))
                workbook = load_workbook(io.BytesIO(archive.read(xlsx[0])), read_only=True)
                try:
                    sheet = workbook.worksheets[0]
                    self.assertEqual(sheet.cell(row=1, column=1).value, "pF - Plot 0")
                    self.assertIsInstance(sheet.cell(row=2, column=1).value, (int, float))
                finally:
                    workbook.close()

    def test_lightweight_runtime_manifest_excludes_classic_statistics_dependencies(self) -> None:
        lines = (PROJECT / "requirements-error-band.txt").read_text(encoding="utf-8").lower()
        self.assertIn("numpy", lines)
        self.assertIn("matplotlib", lines)
        self.assertIn("openpyxl", lines)
        self.assertNotIn("pandas", lines)
        self.assertNotIn("scipy", lines)


if __name__ == "__main__":
    unittest.main()
